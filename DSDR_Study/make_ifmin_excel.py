"""
Excel workbook of the minimum-fault-current calculation (If,min, DG out), from the PowerFactory database:

  1 Method          how the faults were calculated and how If,min is chosen
  2 LG_3ohm_faults  every LG fault through 3 ohm (node, phase): current at the fault, through R1 and R2,
                    and whether the node is in the zone of R1 / R2
  3 If_min_R1_R2    If,min of each recloser zone (values, and the Excel formula for each), the current at the
                    farthest node, the pickup and the check  Ip < If,min
  4 All_sections    If,min of every protected section (Table II of the report)

Input : results/database/5_ShortCircuit_DG_out.csv, results/database/Ifmin_DG_out.csv
Output: results/database/Minimum_Fault_Current_Calculation.xlsx
"""

import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from protection_data import TABLE2_BRANCHES

HERE = os.path.dirname(os.path.abspath(__file__))
DB = os.path.join(HERE, "results", "database")
OUT = os.path.join(DB, "Minimum_Fault_Current_Calculation.xlsx")

ZONES = {"R1": next(b[4] for b in TABLE2_BRANCHES if b[0] == "RG60"),
         "R2": next(b[4] for b in TABLE2_BRANCHES if (b[0], b[1]) == ("632", "671"))}
PICKUP = {"R1": ("CT 900/5, tap 4 A", 720.0, 587.7), "R2": ("CT 1000/5, fast plug 300 A (curve from 600 A)", 600.0, 470.0)}
FARTHEST = {"R1": "652", "R2": "652"}

THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = Font(bold=True)
HFILL = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC")


def grid(ws, r0, rows, widths=None, header=True):
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = ws.cell(row=r0 + i, column=j + 1, value=v)
            if isinstance(v, str) and v.startswith("="):
                c.data_type = "s"                       # show the formula as text, do not evaluate it
            c.border = BOX
            c.alignment = Alignment(vertical="center", wrap_text=True)
            if header and i == 0:
                c.font, c.fill = HEAD, HFILL
    if widths:
        for j, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(j + 1)].width = w


def main():
    sc = [r for r in csv.DictReader(open(os.path.join(DB, "5_ShortCircuit_DG_out.csv")))
          if r["Fault type"] == "LG" and float(r["Zf (ohm)"]) == 3.0]
    wb = Workbook()

    # ---- 1 method ---------------------------------------------------------------------------
    ws = wb.active
    ws.title = "Method"
    lines = [
        ["Minimum fault current If,min of the recloser zones (DG out of service)"],
        [""],
        ["Software", "DIgSILENT PowerFactory 2021, IEEE 13-node feeder model (project IEEE13 DSDR Fuse Coordination)"],
        ["Calculation", "Short circuit, complete method (pre-fault load flow included), initial symmetrical current Ikss"],
        ["Fault", "Line-to-ground (LG) through a fault resistance Rf = 3 ohm"],
        ["Locations", "Every node of the zone and every phase present at that node"],
        ["DG", "Out of service (the pickup is designed without DG)"],
        ["Current read", "Current through the recloser = largest of its three phase currents (fault current + load "
                         "current, added as phasors, because the complete method keeps the load flowing)"],
        ["R1 zone", ", ".join(ZONES["R1"])],
        ["R2 zone", ", ".join(ZONES["R2"]) + "  (downstream of R2)"],
        ["If,min", "Lowest recloser current over all LG 3-ohm faults of the zone  ->  =MIN(IF(...)) in sheet If_min_R1_R2"],
        ["Pickup check", "The pickup must be below If,min, so that the recloser still detects the weakest fault of its zone"],
        [""],
        ["Why If,min is not at the farthest node",
         "The farthest node (652) gives the smallest current AT THE FAULT, but the recloser also carries the load "
         "current of its phase. The smallest RECLOSER current is therefore at 680 on phase B, the phase with the "
         "smallest sum of load and fault current."],
    ]
    for i, row in enumerate(lines, 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if j == 1:
                c.font = HEAD
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 110

    # ---- 2 all LG 3-ohm faults --------------------------------------------------------------
    ws = wb.create_sheet("LG_3ohm_faults")
    rows = [["Faulted node", "Faulted phase", "Rf (ohm)", "Current at the fault (A)", "Current through R1 (A)",
             "Current through R2 (A)", "In R1 zone", "In R2 zone"]]
    for r in sc:
        n = r["Faulted bus"]
        rows.append([n, r["Phases"], 3, float(r["Fault current (A)"]), float(r["R1 current (A)"]),
                     float(r["R2 current (A)"]), "yes" if n in ZONES["R1"] else "no", "yes" if n in ZONES["R2"] else "no"])
    grid(ws, 1, rows, [13, 13, 9, 22, 22, 22, 11, 11])
    last = len(rows)
    for i in range(2, last + 1):
        for col in (4, 5, 6):
            ws.cell(row=i, column=col).number_format = "0.0"
    ws.freeze_panes = "A2"
    # highlight the two minima
    for i in range(2, last + 1):
        if ws.cell(row=i, column=1).value == "680" and ws.cell(row=i, column=2).value == "B":
            for col in (5, 6):
                ws.cell(row=i, column=col).fill = KEY
    ws.cell(row=last + 2, column=1, value="Yellow: the lowest recloser current of each zone (If,min).").font = Font(italic=True)
    ws.cell(row=last + 3, column=1, value="Node 634 is behind the 4.16/0.48 kV transformer: in no recloser zone "
                                          "for this check (its own section is XFM-1).").font = Font(italic=True)

    # ---- 3 If,min: values, and the Excel formula that reproduces each -----------------------
    # (values are written, not only formulas, so the sheet reads correctly even where Excel cannot recalculate)
    ws = wb.create_sheet("If_min_R1_R2")
    D = "LG_3ohm_faults!"
    rng = lambda col: "%s$%s$2:$%s$%d" % (D, col, col, last)
    head = ["Recloser", "Zone (nodes)", "Farthest node", "Current at the fault, farthest node (A)",
            "Recloser current, farthest node (A)", "If,min = lowest recloser current in the zone (A)",
            "Node of If,min", "Phase of If,min", "Inom (A)", "Ip = 1.25 x Inom (A)", "Pickup applied (A)",
            "Setting", "Check: pickup < If,min"]
    grid(ws, 1, [head], [10, 34, 10, 18, 18, 22, 11, 11, 10, 13, 13, 30, 18])
    formulas = [["Column", "Excel formula (R1 row; R2 uses column F / H instead of E / G)"]]
    for k, (dev, key, zone_col, col, zcol) in enumerate((("R1", "Current through R1 (A)", "R1", "E", "G"),
                                                         ("R2", "Current through R2 (A)", "R2", "F", "H"))):
        r = k + 2
        far = FARTHEST[dev]
        setting, ip_applied, inom = PICKUP[dev]
        recs = [x for x in rows[1:] if x[0] in ZONES[dev]]
        cur = 4 if dev == "R1" else 5
        at_far = [x for x in recs if x[0] == far]
        low = min(recs, key=lambda x: x[cur])
        vals = [dev, ", ".join(ZONES[dev]), far, min(x[3] for x in at_far), min(x[cur] for x in at_far), low[cur],
                low[0], low[1], inom, round(1.25 * inom, 1), ip_applied, setting,
                "OK: %.0f < %.1f" % (ip_applied, low[cur]) if ip_applied < low[cur] else "NOT OK"]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BOX
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if c in (4, 5, 6, 10):
                cell.number_format = "0.0"
        ws.cell(row=r, column=6).fill = KEY
        if dev == "R1":
            formulas += [
                ["D  current at the fault, farthest node", "=MIN(IF(%s=C2,%s))   (array formula: Ctrl+Shift+Enter)" % (rng("A"), rng("D"))],
                ["E  recloser current, farthest node", "=MIN(IF(%s=C2,%s))   (array)" % (rng("A"), rng(col))],
                ["F  If,min", "=MIN(IF(%s=\"yes\",%s))   (array)" % (rng(zcol), rng(col))],
                ["G  node of If,min", "=INDEX(%s,MATCH(1,(%s=F2)*(%s=\"yes\"),0))   (array)" % (rng("A"), rng(col), rng(zcol))],
                ["H  phase of If,min", "=INDEX(%s,MATCH(1,(%s=F2)*(%s=\"yes\"),0))   (array)" % (rng("B"), rng(col), rng(zcol))],
                ["J  pickup by eq. (3)", "=1.25*I2"],
                ["M  check", "=IF(K2<F2,\"OK\",\"NOT OK\")"]]
    ws.cell(row=5, column=1, value="How each value is obtained (sheet LG_3ohm_faults holds the data):").font = HEAD
    grid(ws, 6, formulas)
    ws.cell(row=6 + len(formulas) + 1, column=1,
            value="Values from the PowerFactory database: R1 If,min 1074.4 A and R2 If,min 878.5 A, both for the LG "
                  "fault through 3 ohm at 680, phase B.").font = Font(italic=True)

    # ---- 4 all sections ---------------------------------------------------------------------
    ws = wb.create_sheet("All_sections")
    src = list(csv.reader(open(os.path.join(DB, "Ifmin_DG_out.csv"))))
    src[0] = ["Protection section", "Farthest node", "Line length to it (ft)", "Fault type", "Faulted phase",
              "Rf (ohm)", "Current at the fault (A)", "Section current, fault at the farthest node (A)",
              "Lowest LG 3-ohm section current (If,min)"]
    grid(ws, 1, src, [16, 10, 12, 8, 9, 8, 14, 20, 34])
    ws.freeze_panes = "A2"

    try:
        wb.save(OUT)
        out = OUT
    except PermissionError:
        out = OUT.replace(".xlsx", "_new.xlsx")
        wb.save(out)
    print("EXCEL:", out)
    return out


if __name__ == "__main__":
    main()
