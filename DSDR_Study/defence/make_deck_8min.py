"""
8-minute version of the presentation, made from the current deck (which is left unchanged):

  - six detail slides are moved behind the appendix as backup slides, "Thank You" follows the references;
  - the speaker notes are the shortened read-aloud text (defence/deck_8min.json, from make_speaking_script.py);
  - the summary slide gets one bullet with the result of the reference paper;
  - the slide counters are renumbered.

Input : Presentation and report/DSDR_Final_Presentation.pptx
Output: Presentation and report/DSDR_Final_Presentation_8min.pptx
"""

import copy
import json
import os
import re

from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
FOLDER = os.path.join(os.path.dirname(os.path.dirname(HERE)), "Presentation and report")
SRC = os.path.join(FOLDER, "DSDR_Final_Presentation.pptx")
OUT = os.path.join(FOLDER, "DSDR_Final_Presentation_8min.pptx")
PLAN = json.load(open(os.path.join(HERE, "deck_8min.json"), encoding="utf-8"))
ORDER, BACKUP, NOTES = PLAN["order"], PLAN["backup"], PLAN["notes"]

prs = Presentation(SRC)
slides = list(prs.slides)
assert len(slides) == len(ORDER) == 32, (len(slides), len(ORDER))


def text(para):
    return "".join(r.text for r in para.runs)


def set_text(para, new):
    """keep the formatting of the first run"""
    para.runs[0].text = new
    for r in para.runs[1:]:
        r._r.getparent().remove(r._r)


# ---- summary slide: the result of the reference paper next to this study
summary = slides[25 - 1]
done = False
for sh in summary.shapes:
    if sh.has_text_frame and any(text(p).startswith("These counts use the zero-margin") for p in sh.text_frame.paragraphs):
        last = sh.text_frame.paragraphs[-1]
        new = copy.deepcopy(last._p)
        last._p.addnext(new)
        para = sh.text_frame.paragraphs[-1]
        set_text(para, "Reference paper [1]: 30 of 39 with a conventional R2 and 39 of 39 with the DSDR "
                       "(this study: 24 and 38). Same trend; fault levels here follow the IEEE benchmark.")
        done = True
assert done, "summary bullet not added"
for sh in summary.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip() == "5. Results: Summary and Verification":
        set_text(sh.text_frame.paragraphs[0], "5. Results: Summary and Comparison")

# ---- backup slides: say so in the title
for n in BACKUP:
    for sh in slides[n - 1].shapes:
        if sh.has_text_frame and re.match(r"^\d\. (Methodology|Results): ", sh.text_frame.text.strip()):
            para = sh.text_frame.paragraphs[0]
            set_text(para, re.sub(r"^\d\. (Methodology|Results): ", "Backup: ", text(para)))
            break

# ---- speaker notes: the read-aloud text
for n in ORDER:
    tf = slides[n - 1].notes_slide.notes_text_frame
    paras = tf.paragraphs
    if paras and paras[0].runs:
        set_text(paras[0], NOTES[str(n)])
        for extra in paras[1:]:
            extra._p.getparent().remove(extra._p)
    else:
        tf.text = NOTES[str(n)]

# ---- new order
lst = prs.slides._sldIdLst
ids = list(lst)
for el in ids:
    lst.remove(el)
for n in ORDER:
    lst.append(ids[n - 1])

# ---- slide counters "k/n"
for k, slide in enumerate(prs.slides, 1):
    for sh in slide.shapes:
        if sh.has_text_frame and re.fullmatch(r"\d+/\d+", sh.text_frame.text.strip()):
            set_text(sh.text_frame.paragraphs[0], "%d/%d" % (k, len(ORDER)))

prs.save(OUT)
print("saved", OUT)
for k, slide in enumerate(prs.slides, 1):
    title = next((sh.text_frame.text.strip().split("\n")[0] for sh in slide.shapes
                  if sh.has_text_frame and sh.text_frame.text.strip()), "")
    print("%2d  %s" % (k, title[:70]))
