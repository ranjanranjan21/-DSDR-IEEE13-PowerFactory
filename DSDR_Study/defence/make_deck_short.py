"""
Short version of the presentation: only the slides the talk needs, as a separate file. The full deck is
left unchanged. Slides that are not kept are deleted from the short file (they stay in the full deck).

  kept    title, introduction (2), objectives, scope, methodology (4), results (6), conclusions, references,
          thank you, and one appendix slide (the comparison with the reference paper)
  notes   the read-aloud text of make_speaking_script.py
  script  DSDR_Presentation_Script_Short.pdf, numbered as the short deck

Input : Presentation and report/DSDR_Final_Presentation.pptx
Output: Presentation and report/DSDR_Presentation_Short.pptx
"""

import copy
import os
import re

from pptx import Presentation

import make_speaking_script as SS                     # texts per original slide number

HERE = os.path.dirname(os.path.abspath(__file__))
FOLDER = os.path.join(os.path.dirname(os.path.dirname(HERE)), "Presentation and report")
SRC = os.path.join(FOLDER, "DSDR_Final_Presentation.pptx")
OUT = os.path.join(FOLDER, "DSDR_Presentation_Short.pptx")

# original slide numbers, in the order of the short deck
TALK = [1, 3, 4, 5, 6, 7, 8, 10, 12, 16, 18, 21, 23, 24, 25, 26, 27, 32]
APPENDIX = [28]
KEEP = TALK + APPENDIX

prs = Presentation(SRC)
slides = list(prs.slides)
assert len(slides) == 32, len(slides)


def text(para):
    return "".join(r.text for r in para.runs)


def set_text(para, new):
    para.runs[0].text = new
    for r in para.runs[1:]:
        r._r.getparent().remove(r._r)


# ---- summary slide: the result of the reference paper next to this study
summary = slides[25 - 1]
added = False
for sh in summary.shapes:
    if sh.has_text_frame and any(text(p).startswith("These counts use the zero-margin") for p in sh.text_frame.paragraphs):
        last = sh.text_frame.paragraphs[-1]
        last._p.addnext(copy.deepcopy(last._p))
        set_text(sh.text_frame.paragraphs[-1],
                 "Reference paper [1]: 30 of 39 with a conventional R2 and 39 of 39 with the DSDR "
                 "(this study: 24 and 38). Same trend; fault levels here follow the IEEE benchmark.")
        added = True
    if sh.has_text_frame and sh.text_frame.text.strip() == "5. Results: Summary and Verification":
        set_text(sh.text_frame.paragraphs[0], "5. Results: Summary and Comparison")
assert added

# ---- speaker notes
for n in KEEP:
    tf = slides[n - 1].notes_slide.notes_text_frame
    paras = tf.paragraphs
    if paras and paras[0].runs:
        set_text(paras[0], SS.said(n))
        for extra in paras[1:]:
            extra._p.getparent().remove(extra._p)
    else:
        tf.text = SS.said(n)

# ---- keep only KEEP, in that order; the other slides are removed from this file
lst = prs.slides._sldIdLst
ids = list(lst)
for el in ids:
    lst.remove(el)
for n in KEEP:
    lst.append(ids[n - 1])
for n in range(1, 33):
    if n not in KEEP:
        prs.part.drop_rel(ids[n - 1].rId)

# ---- slide counters
for k, slide in enumerate(prs.slides, 1):
    for sh in slide.shapes:
        if sh.has_text_frame and re.fullmatch(r"\d+/\d+", sh.text_frame.text.strip()):
            set_text(sh.text_frame.paragraphs[0], "%d/%d" % (k, len(KEEP)))
prs.save(OUT)

# ---- the script, numbered as the short deck
newno = {old: k + 1 for k, old in enumerate(KEEP)}
main = ["\\slide{%d}{%s}\n" % (newno[n], SS.tex(SS.said(n))) for n in TALK]
app = ["\\slide{%d}{%s}\n" % (newno[n], SS.tex(SS.said(n))) for n in APPENDIX]
nwords = sum(SS.words(SS.TEXT[n][1]) for n in TALK)
doc = SS.DOC.replace("MINUTES", "%.0f" % round(nwords / SS.WPM)).replace("MAIN", "\n".join(main)) \
            .replace("APPENDIX", "\n".join(app))
doc = doc.replace("\\partbar{Backup slides --- only if asked}\n\nBACKUP\n", "") \
         .replace("Appendix slides --- only if asked", "Appendix slide --- only if asked")
assert "BACKUP" not in doc
open(os.path.join(HERE, "DSDR_Presentation_Script_Short.tex"), "w", encoding="utf-8").write(doc)

print("saved", OUT)
print("talk: %d slides, %d words, about %.1f minutes at %d words per minute" % (len(TALK), nwords, nwords / SS.WPM, SS.WPM))
for k, slide in enumerate(prs.slides, 1):
    title = next((sh.text_frame.text.strip().split("\n")[0] for sh in slide.shapes
                  if sh.has_text_frame and sh.text_frame.text.strip()), "")
    print("%2d  %s" % (k, title[:70]))
